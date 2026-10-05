from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_resume_docx(
    resume_profile: dict,
    customized_resume: dict,
    output_path: str
):

    document = Document()

    # --------------------------------
    # PAGE MARGINS
    # --------------------------------

    section = document.sections[0]

    section.top_margin = Pt(40)
    section.bottom_margin = Pt(40)
    section.left_margin = Pt(45)
    section.right_margin = Pt(45)

    # --------------------------------
    # DEFAULT FONT
    # --------------------------------

    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    # --------------------------------
    # NAME
    # --------------------------------

    name = resume_profile.get(
        "name",
        "Candidate"
    )

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run(name)

    run.bold = True
    run.font.size = Pt(20)

    # --------------------------------
    # CONTACT INFORMATION
    # --------------------------------

    contact_items = []

    email = resume_profile.get("email", "")
    phone = resume_profile.get("phone", "")
    location = resume_profile.get("location", "")

    if email:
        contact_items.append(email)

    if phone:
        contact_items.append(phone)

    if location:
        contact_items.append(location)

    if contact_items:

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            " | ".join(contact_items)
        )

        run.font.size = Pt(9)

    # --------------------------------
    # PROFESSIONAL TITLE
    # --------------------------------

    professional_title = customized_resume.get(
        "professional_title",
        ""
    )

    if professional_title:

        paragraph = document.add_paragraph()

        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        run = paragraph.add_run(
            professional_title
        )

        run.bold = True
        run.font.size = Pt(11)

    # --------------------------------
    # SUMMARY
    # --------------------------------

    summary = customized_resume.get(
        "summary",
        ""
    )

    if summary:

        add_section_title(
            document,
            "PROFESSIONAL SUMMARY"
        )

        paragraph = document.add_paragraph()

        paragraph.add_run(summary)

    # --------------------------------
    # SKILLS
    # --------------------------------

    skills = customized_resume.get(
        "skills",
        []
    )

    if skills:

        add_section_title(
            document,
            "SKILLS"
        )

        paragraph = document.add_paragraph()

        paragraph.add_run(
            ", ".join(skills)
        )

    # --------------------------------
    # WORK EXPERIENCE
    # --------------------------------

    work_experience = customized_resume.get(
        "work_experience",
        []
    )

    if work_experience:

        add_section_title(
            document,
            "WORK EXPERIENCE"
        )

        for experience in work_experience:

            company = experience.get(
                "company",
                ""
            )

            position = experience.get(
                "position",
                ""
            )

            duration = experience.get(
                "duration",
                ""
            )

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                position
            )

            run.bold = True

            if company:

                run = paragraph.add_run(
                    f" | {company}"
                )

                run.bold = True

            if duration:

                run = paragraph.add_run(
                    f" | {duration}"
                )

                run.italic = True

            responsibilities = experience.get(
                "responsibilities",
                []
            )

            for responsibility in responsibilities:

                paragraph = document.add_paragraph(
                    style="List Bullet"
                )

                paragraph.add_run(
                    responsibility
                )

    # --------------------------------
    # PROJECTS
    # --------------------------------

    projects = customized_resume.get(
        "projects",
        []
    )

    if projects:

        add_section_title(
            document,
            "PROJECTS"
        )

        for project in projects:

            project_name = project.get(
                "name",
                ""
            )

            technologies = project.get(
                "technologies",
                []
            )

            description = project.get(
                "description",
                ""
            )

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                project_name
            )

            run.bold = True

            if technologies:

                paragraph.add_run(
                    " | " + ", ".join(technologies)
                )

            if description:

                paragraph = document.add_paragraph()

                paragraph.add_run(
                    description
                )

    # --------------------------------
    # EDUCATION
    # --------------------------------

    education = resume_profile.get(
        "education",
        []
    )

    if education:

        add_section_title(
            document,
            "EDUCATION"
        )

        for item in education:

            degree = item.get(
                "degree",
                ""
            )

            institution = item.get(
                "institution",
                ""
            )

            year = item.get(
                "year",
                ""
            )

            paragraph = document.add_paragraph()

            run = paragraph.add_run(
                degree
            )

            run.bold = True

            if institution:

                paragraph.add_run(
                    f" | {institution}"
                )

            if year:

                paragraph.add_run(
                    f" | {year}"
                )

    # --------------------------------
    # CERTIFICATIONS
    # --------------------------------

    certifications = resume_profile.get(
        "certifications",
        []
    )

    if certifications:

        add_section_title(
            document,
            "CERTIFICATIONS"
        )

        for certification in certifications:

            paragraph = document.add_paragraph(
                style="List Bullet"
            )

            paragraph.add_run(
                str(certification)
            )

    # --------------------------------
    # SAVE DOCUMENT
    # --------------------------------

    document.save(output_path)


def add_section_title(
    document,
    title: str
):

    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        title
    )

    run.bold = True
    run.font.size = Pt(11)

    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)